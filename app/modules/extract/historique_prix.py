"""Lecture d'une feuille d'historique de prix observés (module PUR, lot 2 du
31/10/2026, B4).

Une feuille d'un classeur Excel (par exemple « HISTORIQUE POINT P » du
fichier de contrôle d'un fournisseur) liste des prix déjà vus : une ligne par
achat, avec les colonnes « Date », « Facture », « Référence »,
« Désignation », « Quantité », « Unité » et « Prix HT observé ». Ces prix ne
sont **jamais** des prix de grille : ils servent seulement à afficher
l'historique d'un produit (`ingest/historique_prix.py` s'occupe de la base).

Même esprit que `extract/grille_excel.py` (dont les petites fonctions de
lecture de cellule sont réutilisées) : aucune dépendance d'infrastructure,
aucun log, une ligne illisible est comptée dans `erreurs` et jamais
inventée ni devinée. Seules la date, la désignation et le prix sont
obligatoires ; une référence, une quantité, une unité ou un numéro de
facture absents restent vides.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from app.modules.extract.grille_excel import (
    LIGNES_MAX,
    FichierGrilleInvalide,
    _cellule,
    _decimal_depuis_cellule,
    _ligne_vide,
    _lire_reference,
    _normaliser,
    _valeur_texte,
    _valider_signature,
)

_LIGNES_ENTETE_MAX = 10

_TITRES_DATE = {"date", "date facture", "date d observation"}
_TITRES_PIECE = {"facture", "n facture", "numero facture", "piece"}
_TITRES_REFERENCE = {"reference", "ref", "code", "code article"}
_TITRES_DESIGNATION = {"designation", "article", "libelle", "produit"}
_TITRES_QUANTITE = {"quantite", "qte"}
_TITRES_UNITE = {"unite", "u"}
_TITRES_PRIX = {"prix ht observe", "prix ht", "pu ht", "prix observe"}

_FORMAT_FRANCAIS = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$")
_FORMAT_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")


class FeuilleIntrouvable(Exception):
    """La feuille demandée n'existe pas dans le classeur, ou n'a pas les
    colonnes « Date », « Désignation » et « Prix HT observé »."""


@dataclass(frozen=True)
class PrixObserveLu:
    """Une ligne d'historique valide."""

    numero_ligne: int
    date_observation: date
    piece: str | None
    reference: str | None
    designation: str
    quantite: Decimal | None
    unite_brut: str | None
    prix_ht: Decimal


@dataclass(frozen=True)
class ErreurHistorique:
    numero_ligne: int
    message: str


@dataclass(frozen=True)
class HistoriqueLu:
    lignes: tuple[PrixObserveLu, ...]
    erreurs: tuple[ErreurHistorique, ...]


@dataclass
class _Colonnes:
    date: int
    piece: int | None
    reference: int | None
    designation: int
    quantite: int | None
    unite: int | None
    prix: int


def _premier(titres: dict[int, str], reconnus: set[str]) -> int | None:
    return next((i for i, t in titres.items() if t in reconnus), None)


def _colonnes_entete(valeurs: tuple[object, ...]) -> _Colonnes | None:
    titres = {i: _normaliser(v) for i, v in enumerate(valeurs)}
    titres = {i: t for i, t in titres.items() if t}
    col_date = _premier(titres, _TITRES_DATE)
    col_designation = _premier(titres, _TITRES_DESIGNATION)
    col_prix = _premier(titres, _TITRES_PRIX)
    if col_date is None or col_designation is None or col_prix is None:
        return None
    return _Colonnes(
        date=col_date,
        piece=_premier(titres, _TITRES_PIECE),
        reference=_premier(titres, _TITRES_REFERENCE),
        designation=col_designation,
        quantite=_premier(titres, _TITRES_QUANTITE),
        unite=_premier(titres, _TITRES_UNITE),
        prix=col_prix,
    )


def _lire_date(cellule: object) -> date | None:
    """Une date Excel, ou un texte `JJ/MM/AAAA` ou `AAAA-MM-JJ`. `None` si
    la cellule est vide ; `ValueError` si elle n'est pas une date."""
    if cellule is None:
        return None
    if isinstance(cellule, datetime):
        return cellule.date()
    if isinstance(cellule, date):
        return cellule
    texte = str(cellule).strip()
    if not texte:
        return None
    try:
        correspondance = _FORMAT_FRANCAIS.match(texte)
        if correspondance is not None:
            jour, mois, annee = (int(g) for g in correspondance.groups())
            return date(annee, mois, jour)
        correspondance = _FORMAT_ISO.match(texte)
        if correspondance is not None:
            annee, mois, jour = (int(g) for g in correspondance.groups())
            return date(annee, mois, jour)
    except ValueError:
        pass
    raise ValueError(texte)


def _lire_piece(cellule: object) -> str | None:
    return _lire_reference(cellule)


def _traiter_ligne(
    numero: int,
    valeurs: tuple[object, ...],
    colonnes: _Colonnes,
    lignes: list[PrixObserveLu],
    erreurs: list[ErreurHistorique],
) -> None:
    try:
        jour = _lire_date(_cellule(valeurs, colonnes.date))
    except ValueError as exc:
        erreurs.append(ErreurHistorique(numero, f"date illisible : « {exc} »"))
        return
    if jour is None:
        erreurs.append(ErreurHistorique(numero, "date manquante"))
        return

    designation = _valeur_texte(_cellule(valeurs, colonnes.designation))
    if not designation:
        erreurs.append(ErreurHistorique(numero, "désignation manquante"))
        return

    try:
        prix_ht = _decimal_depuis_cellule(_cellule(valeurs, colonnes.prix))
    except ValueError as exc:
        erreurs.append(ErreurHistorique(numero, f"prix illisible : « {exc} »"))
        return
    if prix_ht is None:
        erreurs.append(ErreurHistorique(numero, "prix manquant"))
        return
    if prix_ht <= 0:
        erreurs.append(ErreurHistorique(numero, "le prix doit être supérieur à 0"))
        return

    try:
        quantite = _decimal_depuis_cellule(_cellule(valeurs, colonnes.quantite))
    except ValueError as exc:
        erreurs.append(ErreurHistorique(numero, f"quantité illisible : « {exc} »"))
        return

    lignes.append(
        PrixObserveLu(
            numero_ligne=numero,
            date_observation=jour,
            piece=_lire_piece(_cellule(valeurs, colonnes.piece)),
            reference=_lire_reference(_cellule(valeurs, colonnes.reference)),
            designation=designation,
            quantite=quantite,
            unite_brut=_valeur_texte(_cellule(valeurs, colonnes.unite)) or None,
            prix_ht=prix_ht,
        )
    )


def lire_historique_prix(contenu: bytes, feuille: str) -> HistoriqueLu:
    """Lit la feuille `feuille` du classeur `contenu`. Lève
    `FichierGrilleInvalide` (fichier pas un `.xlsx`, protégé, trop gros) ou
    `FeuilleIntrouvable` (feuille absente ou sans les colonnes attendues)."""
    _valider_signature(contenu)
    try:
        classeur = openpyxl.load_workbook(
            io.BytesIO(contenu), read_only=True, data_only=True
        )
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise FichierGrilleInvalide(
            "le fichier est illisible ou n'est pas un classeur Excel valide (.xlsx)"
        ) from exc

    lignes: list[PrixObserveLu] = []
    erreurs: list[ErreurHistorique] = []
    try:
        if feuille not in classeur.sheetnames:
            raise FeuilleIntrouvable(f"feuille « {feuille} » absente du classeur")
        colonnes: _Colonnes | None = None
        entetes_vus = 0
        explorees = 0
        for ligne_cellules in classeur[feuille].iter_rows():
            explorees += 1
            if explorees > LIGNES_MAX:
                raise FichierGrilleInvalide(
                    f"le fichier dépasse {LIGNES_MAX} lignes utiles"
                )
            numero = next((c.row for c in ligne_cellules if hasattr(c, "row")), None)
            if numero is None:
                continue
            valeurs = tuple(getattr(c, "value", None) for c in ligne_cellules)
            if colonnes is None:
                if entetes_vus >= _LIGNES_ENTETE_MAX:
                    break
                entetes_vus += 1
                colonnes = _colonnes_entete(valeurs)
                continue
            if _ligne_vide(valeurs):
                continue
            _traiter_ligne(numero, valeurs, colonnes, lignes, erreurs)
        if colonnes is None:
            raise FeuilleIntrouvable(
                f"feuille « {feuille} » : colonnes Date, Désignation et "
                "Prix HT observé introuvables"
            )
    finally:
        classeur.close()
    return HistoriqueLu(lignes=tuple(lignes), erreurs=tuple(erreurs))
