"""Tests de `ingest/historique_prix.py` (lot 2 du 31/10/2026, B4) : import
d'une feuille d'historique de prix observés et historique complet d'un
produit. Base de test réelle, données entièrement inventées."""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.extract.historique_prix import FeuilleIntrouvable
from app.modules.ingest.grilles import FournisseurIntrouvable
from app.modules.ingest.historique_prix import (
    FournisseurAmbigu,
    FournisseurNonTrouve,
    historique_complet,
    importer_historique_prix,
)
from app.modules.ingest.models import (
    Document,
    Fournisseur,
    GrillePrix,
    Ligne,
    PrixObserve,
    Utilisateur,
)

pytestmark = [pytest.mark.anyio, pytest.mark.db]

_ENTETE = [
    "Date",
    "Facture",
    "Référence",
    "Désignation",
    "Quantité",
    "Unité",
    "Prix HT observé",
]
_SOURCE = "Fichier de contrôle inventé d'août 2026"


def _classeur(lignes: list[list[object]], feuille: str = "HISTORIQUE X") -> bytes:
    classeur = openpyxl.Workbook()
    page = classeur.active
    assert page is not None
    page.title = feuille
    page.append(_ENTETE)
    for ligne in lignes:
        page.append(ligne)
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


async def _fournisseur(
    db_session: AsyncSession, nom: str = "Négoce Inventé"
) -> Fournisseur:
    fournisseur = Fournisseur(nom=nom, aliases=["NEGOCE INV"])
    db_session.add(fournisseur)
    await db_session.commit()
    await db_session.refresh(fournisseur)
    return fournisseur


async def _observes(db_session: AsyncSession) -> list[PrixObserve]:
    return list((await db_session.execute(select(PrixObserve))).scalars())


async def test_importe_les_lignes_avec_la_cle_de_produit(
    db_session: AsyncSession,
) -> None:
    fournisseur = await _fournisseur(db_session)
    contenu = _classeur(
        [
            [datetime(2026, 8, 3), "F-1", "AB 12", "Enduit", 2, "sac", 13.9],
            [datetime(2026, 8, 5), "F-2", None, "Plaque BA13", 10, "u", 4.5],
            ["n'importe quoi", "F-3", "X", "Y", 1, "u", 1],
        ]
    )

    bilan = await importer_historique_prix(
        db_session,
        fournisseur_nom="négoce inventé",
        contenu=contenu,
        feuille="HISTORIQUE X",
        source=_SOURCE,
    )

    assert (bilan.lues, bilan.importees, bilan.deja_presentes, bilan.erreurs) == (
        2,
        2,
        0,
        1,
    )
    lignes = {ligne.designation: ligne for ligne in await _observes(db_session)}
    assert lignes["Enduit"].cle_produit == "ref:ab12"
    assert lignes["Enduit"].fournisseur_id == fournisseur.id
    assert lignes["Enduit"].piece == "F-1"
    assert lignes["Enduit"].prix_ht == Decimal("13.9")
    assert lignes["Enduit"].source_libelle == _SOURCE
    assert lignes["Plaque BA13"].cle_produit == "nom:plaque ba13"


async def test_import_idempotent(db_session: AsyncSession) -> None:
    await _fournisseur(db_session)
    contenu = _classeur(
        [[datetime(2026, 8, 3), "F-1", "AB 12", "Enduit", 2, "sac", 13.9]]
    )
    for _ in range(2):
        bilan = await importer_historique_prix(
            db_session,
            fournisseur_nom="Négoce Inventé",
            contenu=contenu,
            feuille="HISTORIQUE X",
            source=_SOURCE,
        )

    assert (bilan.importees, bilan.deja_presentes) == (0, 1)
    assert len(await _observes(db_session)) == 1


async def test_ligne_identique_dans_le_meme_fichier_comptee_une_fois(
    db_session: AsyncSession,
) -> None:
    await _fournisseur(db_session)
    ligne = [datetime(2026, 8, 3), "F-1", "AB 12", "Enduit", 2, "sac", 13.9]

    bilan = await importer_historique_prix(
        db_session,
        fournisseur_nom="Négoce Inventé",
        contenu=_classeur([ligne, ligne]),
        feuille="HISTORIQUE X",
        source=_SOURCE,
    )

    assert (bilan.importees, bilan.deja_presentes) == (1, 1)


async def test_simulation_n_ecrit_rien_mais_compte(db_session: AsyncSession) -> None:
    await _fournisseur(db_session)
    contenu = _classeur(
        [[datetime(2026, 8, 3), "F-1", "AB 12", "Enduit", 2, "sac", 13.9]]
    )

    bilan = await importer_historique_prix(
        db_session,
        fournisseur_nom="Négoce Inventé",
        contenu=contenu,
        feuille="HISTORIQUE X",
        source=_SOURCE,
        simulation=True,
    )

    assert bilan.importees == 1
    assert await _observes(db_session) == []


async def test_les_prix_observes_ne_vont_jamais_dans_la_grille(
    db_session: AsyncSession,
) -> None:
    await _fournisseur(db_session)
    await importer_historique_prix(
        db_session,
        fournisseur_nom="Négoce Inventé",
        contenu=_classeur(
            [[datetime(2026, 8, 3), "F-1", "AB 12", "Enduit", 2, "sac", 13.9]]
        ),
        feuille="HISTORIQUE X",
        source=_SOURCE,
    )

    assert (await db_session.execute(select(GrillePrix))).scalars().all() == []


async def test_fournisseur_inconnu_ou_ambigu_et_feuille_absente(
    db_session: AsyncSession,
) -> None:
    await _fournisseur(db_session, "Dupont Matériaux")
    await _fournisseur(db_session, "Dupont  Materiaux")
    contenu = _classeur([])

    with pytest.raises(FournisseurNonTrouve):
        await importer_historique_prix(
            db_session,
            fournisseur_nom="Inconnu",
            contenu=contenu,
            feuille="HISTORIQUE X",
            source=_SOURCE,
        )
    with pytest.raises(FournisseurAmbigu):
        await importer_historique_prix(
            db_session,
            fournisseur_nom="dupont materiaux",
            contenu=contenu,
            feuille="HISTORIQUE X",
            source=_SOURCE,
        )
    await _fournisseur(db_session, "Seul Fournisseur")
    with pytest.raises(FeuilleIntrouvable):
        await importer_historique_prix(
            db_session,
            fournisseur_nom="Seul Fournisseur",
            contenu=contenu,
            feuille="ABSENTE",
            source=_SOURCE,
        )


# --------------------------------------------------------------------------
# Historique complet d'un produit
# --------------------------------------------------------------------------


async def _grille(
    db_session: AsyncSession,
    fournisseur: Fournisseur,
    *,
    date_effet: date,
    prix: str,
    source: str,
    remise: str | None = None,
    **extra: object,
) -> GrillePrix:
    prix_net = Decimal(prix) * (Decimal(100) - Decimal(remise or 0)) / Decimal(100)
    ligne = GrillePrix(
        fournisseur_id=fournisseur.id,
        cle_produit="ref:ab12",
        reference="AB 12",
        designation="Enduit",
        unite_brut="sac",
        prix_ht=Decimal(prix),
        remise_pct=Decimal(remise) if remise is not None else None,
        prix_net=prix_net,
        date_effet=date_effet,
        source=source,
        **extra,
    )
    db_session.add(ligne)
    await db_session.commit()
    return ligne


async def test_historique_complet_melange_les_trois_sources_par_date_decroissante(
    db_session: AsyncSession,
) -> None:
    fournisseur = await _fournisseur(db_session)
    auteur = Utilisateur(
        email=f"a.{uuid.uuid4()}@example.invalid",
        role="admin",
        is_active=True,
        nom="Alice",
    )
    db_session.add(auteur)
    bon = Document(
        type="bl",
        source="telegram",
        statut="recu",
        fichier_path="chemin-invente.jpg",
        numero="BL-77",
        fournisseur_id=fournisseur.id,
        date_doc=date(2026, 8, 20),
    )
    db_session.add(bon)
    await db_session.commit()
    await _grille(
        db_session,
        fournisseur,
        date_effet=date(2026, 1, 1),
        prix="10.00",
        source="saisie",
        auteur_id=auteur.id,
    )
    await _grille(
        db_session,
        fournisseur,
        date_effet=date(2026, 8, 25),
        prix="12.00",
        remise="10",
        source="validation",
        document_id=bon.id,
    )
    # Prix vus sur un bon du même fournisseur : code lu = référence du produit.
    db_session.add(
        Ligne(
            document_id=bon.id,
            code_brut="ab-12",
            article_brut="Enduit",
            pu_ht=Decimal("11.50"),
            qte_livree=Decimal("3"),
            type_ligne="article",
            ordre=0,
        )
    )
    # Autre code : n'apparaît pas.
    db_session.add(
        Ligne(
            document_id=bon.id,
            code_brut="ZZ99",
            pu_ht=Decimal("1.00"),
            type_ligne="article",
            ordre=1,
        )
    )
    db_session.add(
        PrixObserve(
            fournisseur_id=fournisseur.id,
            cle_produit="ref:ab12",
            reference="AB 12",
            designation="Enduit",
            date_observation=date(2026, 7, 1),
            prix_ht=Decimal("9.80"),
            quantite=Decimal("4"),
            piece="F-9",
            source_libelle=_SOURCE,
        )
    )
    await db_session.commit()

    entrees = await historique_complet(db_session, fournisseur.id, "ref:ab12")

    assert [(e.type, e.date) for e in entrees] == [
        ("grille", date(2026, 8, 25)),
        ("document", date(2026, 8, 20)),
        ("observe", date(2026, 7, 1)),
        ("grille", date(2026, 1, 1)),
    ]
    grille_recente, vu, observe, grille_ancienne = entrees
    assert grille_recente.origine == "Validé depuis le bon n° BL-77"
    assert grille_recente.prix_ht == Decimal("12.00")
    assert grille_recente.remise_pct == Decimal("10")
    assert grille_recente.prix_net == Decimal("10.8")
    assert vu.numero_document == "BL-77"
    assert vu.prix_ht == Decimal("11.50")
    assert vu.quantite == Decimal("3")
    assert vu.document_id == bon.id
    assert observe.origine == _SOURCE
    assert observe.numero_document == "F-9"
    assert grille_ancienne.origine == "Saisi par Alice"


async def test_historique_complet_ignore_documents_archives_et_autres_fournisseurs(
    db_session: AsyncSession,
) -> None:
    fournisseur = await _fournisseur(db_session)
    autre = await _fournisseur(db_session, "Autre Négoce")
    await _grille(
        db_session, fournisseur, date_effet=date(2026, 1, 1), prix="10", source="saisie"
    )
    for proprietaire, statut in ((fournisseur, "archive"), (autre, "recu")):
        doc = Document(
            type="facture",
            source="upload",
            statut=statut,
            fichier_path=f"c-{uuid.uuid4()}.pdf",
            numero=f"N-{uuid.uuid4()}",
            fournisseur_id=proprietaire.id,
            date_doc=date(2026, 8, 1),
        )
        db_session.add(doc)
        await db_session.commit()
        db_session.add(
            Ligne(
                document_id=doc.id,
                code_brut="AB12",
                pu_ht=Decimal("5"),
                type_ligne="article",
                ordre=0,
            )
        )
    await db_session.commit()

    entrees = await historique_complet(db_session, fournisseur.id, "ref:ab12")

    assert [e.type for e in entrees] == ["grille"]


async def test_historique_complet_limite_a_100_lignes(db_session: AsyncSession) -> None:
    fournisseur = await _fournisseur(db_session)
    db_session.add_all(
        PrixObserve(
            fournisseur_id=fournisseur.id,
            cle_produit="nom:x",
            designation="X",
            date_observation=date(2025, 1, 1)
            + (date(2025, 1, 2) - date(2025, 1, 1)) * i,
            prix_ht=Decimal("1"),
            source_libelle=_SOURCE,
        )
        for i in range(120)
    )
    await db_session.commit()

    entrees = await historique_complet(db_session, fournisseur.id, "nom:x")

    assert len(entrees) == 100
    assert entrees[0].date > entrees[-1].date


async def test_historique_complet_fournisseur_introuvable(
    db_session: AsyncSession,
) -> None:
    with pytest.raises(FournisseurIntrouvable):
        await historique_complet(db_session, uuid.uuid4(), "ref:x")
